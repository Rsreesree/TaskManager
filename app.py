# app.py - Flask Backend with Excel Storage + Auto Browser Open
from flask import Flask, render_template, request, jsonify, redirect, url_for, session
from datetime import datetime, date
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
import os
import smtplib
from email.message import EmailMessage
import webbrowser
import threading
import time
import requests
import hashlib

app = Flask(__name__, template_folder="templates", static_folder="static")
app.secret_key = "supersecretkey123!@#"

TASKS_FILE = 'tasks.xlsx'
USERS_FILE = 'users.xlsx'


# --------------------- Excel Initialization ---------------------
def init_excel():
    if not os.path.exists(TASKS_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tasks"
        headers = ['ID', 'Task', 'Category', 'Priority', 'Assigned To', 'Contact', 'Due Date', 'Status', 'Created', 'Completed']
        ws.append(headers)
        # Style header row
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        # Set column widths
        widths = [15, 40, 12, 12, 20, 25, 15, 12, 20, 20]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        wb.save(TASKS_FILE)
        print(f"✓ Excel file '{TASKS_FILE}' created successfully!")

# --------------------- Excel Utilities ---------------------
def get_all_tasks():
    init_excel()
    wb = openpyxl.load_workbook(TASKS_FILE)
    ws = wb.active
    tasks = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            # Parse due date
            try:
                due_date_str = row[6].strftime('%Y-%m-%d') if isinstance(row[6], (datetime, date)) else str(row[6]) if row[6] else None
            except:
                due_date_str = None
            # Parse created & completed
            try:
                created_at_str = row[8].strftime('%Y-%m-%d %H:%M:%S') if isinstance(row[8], datetime) else str(row[8]) if row[8] else None
                completed_at_str = row[9].strftime('%Y-%m-%d %H:%M:%S') if isinstance(row[9], datetime) else str(row[9]) if row[9] else None
            except:
                created_at_str = completed_at_str = None
            tasks.append({
                'id': int(row[0]),
                'text': str(row[1]) if row[1] else '',
                'category': str(row[2]) if row[2] else 'other',
                'priority': str(row[3]) if row[3] else 'medium',
                'assignedTo': str(row[4]) if row[4] else '',
                'contact': str(row[5]) if row[5] else '',
                'dueDate': due_date_str,
                'completed': str(row[7]).lower() == 'completed',
                'createdAt': created_at_str,
                'completedAt': completed_at_str
            })
    wb.close()
    return tasks

def add_task_to_excel(task_id, text, category, priority, assigned_to, contact, due_date):
    wb = openpyxl.load_workbook(TASKS_FILE)
    ws = wb.active
    due_date_obj = None
    if due_date:
        try:
            due_date_obj = datetime.strptime(due_date, '%Y-%m-%d')
        except:
            due_date_obj = None
    ws.append([task_id, text, category, priority, assigned_to, contact, due_date_obj, 'active', datetime.now(), None])
    # Style row based on priority
    row_num = ws.max_row
    fill_color = {"high":"FFE6E6","medium":"FFF4E6","low":"E6F4EA"}.get(priority,"E6F4EA")
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    for cell in ws[row_num]:
        cell.fill = fill
        cell.alignment = Alignment(horizontal='left', vertical='center')
    wb.save(TASKS_FILE)
    wb.close()
    print(f"✓ Task added: {text}")

def update_task_status(task_id, completed):
    wb = openpyxl.load_workbook(TASKS_FILE)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value == task_id:
            row[7].value = 'completed' if completed else 'active'
            row[9].value = datetime.now() if completed else None
            break
    wb.save(TASKS_FILE)
    wb.close()

def delete_task_from_excel(task_id):
    wb = openpyxl.load_workbook(TASKS_FILE)
    ws = wb.active
    found = False
    for row in ws.iter_rows(min_row=2):
        if row[0].value == task_id:
            ws.delete_rows(row[0].row, 1)
            found = True
            break
    wb.save(TASKS_FILE)
    wb.close()
    return found

# --------------------- Email Messaging ---------------------
def send_email(to, subject, body):
    msg = EmailMessage()
    msg['Subject'] = subject
    msg['From'] = 'your_email@gmail.com'  # Replace
    msg['To'] = to
    msg.set_content(body)
    with smtplib.SMTP_SSL('smtp.gmail.com', 465) as smtp:
        smtp.login('your_email@gmail.com', 'your_app_password')  # Replace with App Password
        smtp.send_message(msg)

@app.route('/api/send_message', methods=['POST'])
def send_message():
    data = request.json
    to = data.get('to')
    message = data.get('message')
    subject = data.get('subject','Task Manager Notification')
    try:
        send_email(to, subject, message)
        return {'success': True, 'message': 'Message sent successfully'}
    except Exception as e:
        return {'success': False, 'message': str(e)}, 500
    
def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

# ---------------- Initialize Users Sheet ----------------
def init_users_excel():
    if not os.path.exists(USERS_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Users'
        ws.append(['Username', 'PasswordHash'])
        # Create default admin
        ws.append(['admin', hash_password('admin123')])
        wb.save(USERS_FILE)
        print("✓ Users Excel created with default admin")


# ---------------- Add User ----------------
def add_user(username, password):
    if not os.path.exists(USERS_FILE):
        init_users_excel()
    wb = openpyxl.load_workbook(USERS_FILE)
    ws = wb['Users']
    # Check for duplicate
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == username:
            wb.close()
            return False
    ws.append([username, hash_password(password)])
    wb.save(USERS_FILE)
    wb.close()
    return True


# ---------------- Verify User ----------------
def verify_user(username, password):
    if not os.path.exists(USERS_FILE):
        return False
    wb = openpyxl.load_workbook(USERS_FILE)
    if 'Users' not in wb.sheetnames:
        wb.close()
        return False
    ws = wb['Users']
    pw_hash = hash_password(password)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == username and row[1] == pw_hash:
            wb.close()
            return True
    wb.close()
    return False


# --------------------- Flask Routes ---------------------
@app.route("/", methods=['GET', 'POST'])
def login():
    init_users_excel()
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        if verify_user(username, password):
            session['username'] = username
            return redirect(url_for('index'))
        else:
            return render_template('login.html', error="Invalid username or password")
    return render_template("login.html")

@app.route("/register", methods=['GET', 'POST'])
def register():
    init_users_excel()
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        wb = openpyxl.load_workbook(TASKS_FILE)
        ws = wb['Users']
        # Check if username exists
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == username:
                wb.close()
                return render_template('register.html', error="Username already exists")
        # Add new user
        ws.append([username, hash_password(password)])
        wb.save(TASKS_FILE)
        wb.close()
        return redirect(url_for('login'))
    return render_template('register.html')


@app.route("/index")
def index():
    if 'username' not in session:
        return redirect(url_for('login'))
    tasks = get_all_tasks()
    return render_template('index.html', tasks=tasks, username=session['username'])


@app.route("/api/signup", methods=['POST'])
def api_signup():
    init_users_excel()
    data = request.get_json()
    username = data.get('username')
    password = data.get('password')

    if not username or not password:
        return jsonify({'success': False, 'message': 'Username and password required'}), 400

    if add_user(username, password):
        return jsonify({'success': True, 'message': 'Account created successfully'}), 200
    else:
        return jsonify({'success': False, 'message': 'Username already exists'}), 400

@app.route("/api/login", methods=['POST'])
def api_login():
    init_users_excel()
    data = request.get_json()
    username = data.get('username')
    password = data.get('password')

    if verify_user(username, password):
        session['username'] = username
        return jsonify({'success': True, 'message': 'Login successful'}), 200
    else:
        return jsonify({'success': False, 'message': 'Invalid username or password'}), 401


@app.route('/api/tasks', methods=['GET'])
def get_tasks():
    tasks = get_all_tasks()
    date_filter = request.args.get('date')
    person_filter = request.args.get('person')
    priority_filter = request.args.get('priority')
    if date_filter:
        tasks = [t for t in tasks if t['dueDate'] == date_filter]
    if person_filter and person_filter != 'all':
        tasks = [t for t in tasks if t['assignedTo'] == person_filter]
    if priority_filter and priority_filter != 'all':
        tasks = [t for t in tasks if t['priority'] == priority_filter]
    return jsonify(tasks)

@app.route('/api/persons', methods=['GET'])
def get_persons():
    tasks = get_all_tasks()
    persons = set()
    for task in tasks:
        if task['assignedTo'].strip():
            persons.add(task['assignedTo'].strip())
    return jsonify(sorted(list(persons)))

@app.route('/api/tasks', methods=['POST'])
def create_task():
    data = request.json
    task_id = int(datetime.now().timestamp()*1000)
    add_task_to_excel(
        task_id,
        data.get('text',''),
        data.get('category','personal'),
        data.get('priority','medium'),
        data.get('assignedTo',''),
        data.get('contact',''),
        data.get('dueDate','')
    )
    return jsonify({
        'id': task_id,
        'text': data.get('text',''),
        'category': data.get('category','personal'),
        'priority': data.get('priority','medium'),
        'assignedTo': data.get('assignedTo',''),
        'contact': data.get('contact',''),
        'dueDate': data.get('dueDate',''),
        'completed': False,
        'createdAt': datetime.now().isoformat()
    }),201

@app.route('/api/tasks/<int:task_id>', methods=['PUT'])
def update_task(task_id):
    data = request.json
    update_task_status(task_id, data.get('completed', False))
    return jsonify({'success': True, 'message': 'Task updated successfully'})

@app.route('/api/tasks/<int:task_id>', methods=['DELETE'])
def delete_task(task_id):
    success = delete_task_from_excel(task_id)
    if success:
        return jsonify({'success': True, 'message': 'Task deleted successfully'})
    return jsonify({'success': False, 'message': 'Task not found'}),404

@app.route('/api/stats', methods=['GET'])
def get_stats():
    tasks = get_all_tasks()
    today = datetime.now().date()
    stats = {
        'total': len(tasks),
        'completed': len([t for t in tasks if t['completed']]),
        'active': len([t for t in tasks if not t['completed']]),
        'high_priority': len([t for t in tasks if t['priority']=='high' and not t['completed']]),
        'overdue': sum(1 for t in tasks if not t['completed'] and t['dueDate'] and datetime.strptime(t['dueDate'],'%Y-%m-%d').date()<today)
    }
    return jsonify(stats)

# --------------------- Auto Open Browser ---------------------
def open_browser():
    time.sleep(1)  # wait for Flask server to start
    webbrowser.open_new("http://127.0.0.1:5000/")

if __name__ == "__main__":
    threading.Thread(target=open_browser).start()
    app.run(debug=False)



